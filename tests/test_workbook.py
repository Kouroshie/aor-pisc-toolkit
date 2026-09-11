"""The Excel workbook: every input in one file, and back again.

A YAML project file is the better record. The workbook is the one a team will
actually fill in, because the people holding the reservoir data work in
spreadsheets. It is only worth having if nothing is lost in the round trip,
which is what these tests are for.
"""

from __future__ import annotations

import os

import pytest

from containment.config import Project
from containment.io import workbook

pytest.importorskip("openpyxl")

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOWCASE = os.path.join(REPO, "examples", "showcase_full.yaml")


@pytest.fixture
def book(tmp_path):
    os.chdir(REPO)                       # the penetrations path is repo-relative
    project = Project.from_yaml(SHOWCASE)
    path = str(tmp_path / "inputs.xlsx")
    workbook.to_excel(project, path)
    return project, path


def test_every_sheet_is_written(book):
    from openpyxl import load_workbook

    _, path = book
    names = load_workbook(path).sheetnames
    for sheet in ("Read me", "Project", "Units", "Injection zones", "USDW",
                  "Relative permeability", "Threshold", "Model", "Plume",
                  "Wells", "Faults", "Penetrations", "Uncertainty"):
        assert sheet in names, sheet


def test_the_round_trip_preserves_the_project(book):
    original, path = book
    back = Project.from_dict(workbook.from_excel(path))

    assert back.name == original.name
    assert back.start_date == original.start_date
    assert [z.name for z in back.zones] == [z.name for z in original.zones]
    assert len(back.wells) == len(original.wells)
    assert back.total_injected_mass() == pytest.approx(
        original.total_injected_mass(), rel=1e-9)
    assert len(back.faults) == len(original.faults)
    assert back.end_time == pytest.approx(original.end_time)
    assert back.plume_cutoff == original.plume_cutoff
    assert back.uncertainty.get("realisations") == \
        original.uncertainty.get("realisations")


def test_zone_properties_survive_in_si(book):
    original, path = book
    back = Project.from_dict(workbook.from_excel(path))
    for a, b in zip(original.zones, back.zones, strict=True):
        assert b.top_depth == pytest.approx(a.top_depth)
        assert b.thickness == pytest.approx(a.thickness)
        assert b.permeability == pytest.approx(a.permeability)
        assert b.initial_pressure == pytest.approx(a.initial_pressure)
        assert b.porosity == pytest.approx(a.porosity)
    # each zone keeps its own confining interval
    assert back.seals[0].top_depth != back.seals[1].top_depth


def test_the_penetration_list_travels_inside_the_workbook(book):
    from containment import corrective, workflow

    original, path = book
    back = Project.from_dict(workbook.from_excel(path))

    a = corrective.load_wells_csv(original.penetrations_csv,
                                  unit=original.penetrations_unit,
                                  crs=workflow.project_crs(original))
    b = corrective.load_wells_csv(back.penetrations_csv,
                                  unit=back.penetrations_unit,
                                  crs=workflow.project_crs(back))
    assert [w.name for w in a] == [w.name for w in b]
    assert a[0].total_depth == pytest.approx(b[0].total_depth)


def test_a_stepped_schedule_survives(book):
    original, path = book
    back = Project.from_dict(workbook.from_excel(path))
    stepped = [w for w in original.wells if len(w.schedule) > 2]
    assert stepped, "the showcase should carry a stepped well"
    same = [w for w in back.wells if w.name == stepped[0].name][0]
    assert len(same.schedule) == len(stepped[0].schedule)
    assert same.total_mass() == pytest.approx(stepped[0].total_mass(), rel=1e-9)


def test_a_single_zone_project_round_trips_as_one_zone(tmp_path):
    d = {
        "project": {"name": "one zone"},
        "units": {"length": "ft", "depth": "ft", "pressure": "psi",
                  "temperature": "F", "permeability": "mD", "rate": "MMT/yr",
                  "time": "yr", "compressibility": "1/psi"},
        "formation": {
            "injection_zone": {"top_depth": 6000, "thickness": 250,
                               "porosity": 0.18, "permeability": 150,
                               "temperature": 150, "salinity_ppm": 90000,
                               "initial_pressure": 2600},
            "confining_zone": {"top_depth": 5700, "base_depth": 6000},
            "usdw": {"base_depth": 1200, "initial_pressure": 520,
                     "temperature": 80, "salinity_ppm": 800}},
        "wells": [{"name": "INJ-1", "x": 0, "y": 0, "rate": 0.5,
                   "start_year": 0, "stop_year": 10}],
    }
    path = str(tmp_path / "one.xlsx")
    workbook.to_excel(Project.from_dict(d), path)
    back = Project.from_dict(workbook.from_excel(path))
    assert not back.is_stacked
    assert back.injection_zone.top_depth == pytest.approx(
        Project.from_dict(d).injection_zone.top_depth)
    assert back.confining_zone.top_depth == pytest.approx(
        Project.from_dict(d).confining_zone.top_depth)


def test_load_sniffs_the_format(book):
    _, path = book
    assert Project.load(path).name == Project.load(SHOWCASE).name
