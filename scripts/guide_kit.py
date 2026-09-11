"""Shared helpers for building the Containment user guide.

Holds the document object, the run data and the formatting primitives, so
``build_guide.py`` reads as the guide's outline rather than as Word plumbing.
"""

from __future__ import annotations

import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = (sys.argv[1] if len(sys.argv) > 1
       else os.path.join(os.path.expanduser("~"), "Documents", "Containment_Guide"))
FIG = os.path.join(OUT, "figures")
DATA = json.load(open(os.path.join(OUT, "data.json"), encoding="utf-8"))

INK = RGBColor(0x0B, 0x0B, 0x0B)
BRAND = RGBColor(0x18, 0x4F, 0x95)
MUTED = RGBColor(0x52, 0x51, 0x4E)

doc = Document()

# --------------------------------------------------------------- styling
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, colour, space_before in (
        ("Heading 1", 20, BRAND, 22), ("Heading 2", 15, BRAND, 18),
        ("Heading 3", 12, INK, 14), ("Heading 4", 11, MUTED, 10)):
    st = doc.styles[name]
    st.font.name = "Georgia"
    st.font.size = Pt(size)
    st.font.color.rgb = colour
    st.font.bold = True
    st.paragraph_format.space_before = Pt(space_before)
    st.paragraph_format.space_after = Pt(6)


def h(level: int, text: str):
    return doc.add_heading(text, level=level)


def p(text: str = "", *, italic=False, bold=False, size=None, colour=None,
      align=None, space_after=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.italic, run.bold = italic, bold
    if size:
        run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = colour
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def rich(*parts):
    """A paragraph from (text, bold, italic) triples."""
    par = doc.add_paragraph()
    for item in parts:
        text, bold, italic = (item + (False, False))[:3] if isinstance(item, tuple) \
            else (item, False, False)
        run = par.add_run(text)
        run.bold, run.italic = bold, italic
    return par


def bullet(text: str, level: int = 0):
    par = doc.add_paragraph(text, style="List Bullet")
    par.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    par.paragraph_format.space_after = Pt(4)
    return par


def numbered(text: str):
    par = doc.add_paragraph(text, style="List Number")
    par.paragraph_format.space_after = Pt(4)
    return par


def shade(cell, hex_colour: str):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def code(text: str, caption: str = ""):
    if caption:
        p(caption, italic=True, size=9, colour=MUTED, space_after=2)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, "F4F3F0")
    cell.text = ""
    for i, line in enumerate(text.rstrip().splitlines()):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        run = par.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def table(rows: list[dict], columns=None, caption: str = "", max_rows: int = 40,
          widths=None):
    if not rows:
        p("(nothing to show for this run)", italic=True, colour=MUTED)
        return None
    cols = columns or list(rows[0].keys())
    if caption:
        p(caption, italic=True, size=9, colour=MUTED, space_after=3)
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]
        cell.text = str(c).replace("_", " ")
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(8.5)
    for row in rows[:max_rows]:
        cells = t.add_row().cells
        for i, c in enumerate(cols):
            v = row.get(c, "")
            if isinstance(v, float):
                v = (f"{v:,.3g}" if abs(v) < 0.01 and v != 0 else f"{v:,.2f}")
            cells[i].text = str(v)
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.5)
    if len(rows) > max_rows:
        p(f"({len(rows) - max_rows} further rows omitted; the full table is in "
          "the CSV and HTML exports)", italic=True, size=9, colour=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name: str, caption: str, width: float = 6.2):
    path = os.path.join(FIG, name + ".png")
    if not os.path.exists(path):
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par = p(caption, italic=True, size=9, colour=MUTED,
            align=WD_ALIGN_PARAGRAPH.CENTER)
    par.paragraph_format.space_after = Pt(12)


def callout(title: str, body: str, fill: str = "EAF2FB"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    par = cell.paragraphs[0]
    run = par.add_run(title)
    run.bold = True
    run.font.size = Pt(10)
    par2 = cell.add_paragraph()
    run2 = par2.add_run(body)
    run2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def num(x, fmt="{:,.0f}", default="n/a"):
    try:
        if x is None:
            return default
        return fmt.format(float(x))
    except (TypeError, ValueError):
        return str(x)


D = DATA
PROJ, AOR = D["project"], D["aor"]
SUM = D["summary"]
